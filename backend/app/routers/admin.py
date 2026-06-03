import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import re
from zipfile import BadZipFile

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Distribution, DonationReceived, Household, User
from app.schemas import (
    DistributionCreate,
    DistributionOut,
    DistributionUpdate,
    DonationCreate,
    DonationOut,
    DonationUpdate,
    HouseholdCreate,
    HouseholdOut,
    HouseholdUpdate,
    ImportErrorRow,
    ImportPreview,
    PublicSummary,
    UserOut,
)
from app.security import get_current_admin
from app.services import calculate_summary


router = APIRouter(prefix="/admin", tags=["admin"], dependencies=[Depends(get_current_admin)])


@router.get("/me", response_model=UserOut)
def admin_me(user: User = Depends(get_current_admin)) -> User:
    return user


@router.get("/summary", response_model=PublicSummary)
def admin_summary(db: Session = Depends(get_db)) -> PublicSummary:
    return calculate_summary(db)


def commit_or_conflict(db: Session) -> None:
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="A record with this unique value already exists") from exc


def serialize_distribution(row: Distribution) -> DistributionOut:
    return DistributionOut(
        id=row.id,
        distribution_code=row.distribution_code,
        household_id=row.household_id,
        household_code=row.household.household_code,
        location=row.household.location,
        amount=row.amount,
        distribution_date=row.distribution_date,
        assistance_type=row.assistance_type,
        notes=row.notes,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def resolve_household_id(db: Session, household_id: int | None, household_code: str | None) -> int:
    if household_id:
        household = db.get(Household, household_id)
    elif household_code:
        household = db.query(Household).filter(Household.household_code == household_code).first()
    else:
        raise HTTPException(status_code=422, detail="household_id or household_code is required")
    if not household:
        raise HTTPException(status_code=404, detail="Household not found")
    return household.id


@router.get("/donations", response_model=list[DonationOut])
def list_donations(db: Session = Depends(get_db)) -> list[DonationReceived]:
    return db.query(DonationReceived).order_by(DonationReceived.received_date.desc(), DonationReceived.id.desc()).all()


@router.post("/donations", response_model=DonationOut, status_code=status.HTTP_201_CREATED)
def create_donation(payload: DonationCreate, db: Session = Depends(get_db)) -> DonationReceived:
    data = payload.model_dump()
    data["currency"] = "USD"
    row = DonationReceived(**data)
    db.add(row)
    commit_or_conflict(db)
    db.refresh(row)
    return row


@router.put("/donations/{donation_id}", response_model=DonationOut)
def update_donation(donation_id: int, payload: DonationUpdate, db: Session = Depends(get_db)) -> DonationReceived:
    row = db.get(DonationReceived, donation_id)
    if not row:
        raise HTTPException(status_code=404, detail="Donation not found")
    data = payload.model_dump()
    data["currency"] = "USD"
    for key, value in data.items():
        setattr(row, key, value)
    commit_or_conflict(db)
    db.refresh(row)
    return row


@router.delete("/donations/{donation_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_donation(donation_id: int, db: Session = Depends(get_db)) -> None:
    row = db.get(DonationReceived, donation_id)
    if not row:
        raise HTTPException(status_code=404, detail="Donation not found")
    db.delete(row)
    db.commit()


@router.get("/households", response_model=list[HouseholdOut])
def list_households(db: Session = Depends(get_db)) -> list[Household]:
    return db.query(Household).order_by(Household.household_code.asc()).all()


@router.post("/households", response_model=HouseholdOut, status_code=status.HTTP_201_CREATED)
def create_household(payload: HouseholdCreate, db: Session = Depends(get_db)) -> Household:
    row = Household(**payload.model_dump())
    db.add(row)
    commit_or_conflict(db)
    db.refresh(row)
    return row


@router.put("/households/{household_id}", response_model=HouseholdOut)
def update_household(household_id: int, payload: HouseholdUpdate, db: Session = Depends(get_db)) -> Household:
    row = db.get(Household, household_id)
    if not row:
        raise HTTPException(status_code=404, detail="Household not found")
    for key, value in payload.model_dump().items():
        setattr(row, key, value)
    commit_or_conflict(db)
    db.refresh(row)
    return row


@router.delete("/households/{household_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_household(household_id: int, db: Session = Depends(get_db)) -> None:
    row = db.get(Household, household_id)
    if not row:
        raise HTTPException(status_code=404, detail="Household not found")
    db.delete(row)
    db.commit()


@router.get("/distributions", response_model=list[DistributionOut])
def list_distributions(db: Session = Depends(get_db)) -> list[DistributionOut]:
    rows = db.query(Distribution).join(Distribution.household).order_by(Distribution.distribution_date.desc()).all()
    return [serialize_distribution(row) for row in rows]


@router.post("/distributions", response_model=DistributionOut, status_code=status.HTTP_201_CREATED)
def create_distribution(payload: DistributionCreate, db: Session = Depends(get_db)) -> DistributionOut:
    data = payload.model_dump()
    data["household_id"] = resolve_household_id(db, data.pop("household_id"), data.pop("household_code"))
    data["currency"] = "USD"
    row = Distribution(**data)
    db.add(row)
    commit_or_conflict(db)
    db.refresh(row)
    return serialize_distribution(row)


@router.put("/distributions/{distribution_id}", response_model=DistributionOut)
def update_distribution(distribution_id: int, payload: DistributionUpdate, db: Session = Depends(get_db)) -> DistributionOut:
    row = db.get(Distribution, distribution_id)
    if not row:
        raise HTTPException(status_code=404, detail="Distribution not found")
    data = payload.model_dump()
    data["household_id"] = resolve_household_id(db, data.pop("household_id"), data.pop("household_code"))
    data["currency"] = "USD"
    for key, value in data.items():
        setattr(row, key, value)
    commit_or_conflict(db)
    db.refresh(row)
    return serialize_distribution(row)


@router.delete("/distributions/{distribution_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_distribution(distribution_id: int, db: Session = Depends(get_db)) -> None:
    row = db.get(Distribution, distribution_id)
    if not row:
        raise HTTPException(status_code=404, detail="Distribution not found")
    db.delete(row)
    db.commit()


def parse_decimal(value: str, field: str, errors: list[str]) -> Decimal:
    try:
        amount = Decimal(value)
        if amount <= 0:
            errors.append(f"{field} must be greater than 0")
        return amount
    except (InvalidOperation, TypeError):
        errors.append(f"{field} must be a valid decimal")
        return Decimal("0")


def parse_date(value: object, field: str, errors: list[str]) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    normalized = (
        str(value or "")
        .strip()
        .replace("\ufeff", "")
        .replace("\u200b", "")
        .replace("\u2010", "-")
        .replace("\u2011", "-")
        .replace("\u2012", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
    )
    match = re.match(r"^(\d{4}-\d{2}-\d{2})(?:[T\s].*)?$", normalized)
    try:
        return date.fromisoformat(match.group(1) if match else normalized)
    except ValueError:
        errors.append(f"{field} must use YYYY-MM-DD")
        return date.today()


async def read_csv(upload: UploadFile) -> list[dict[str, str]]:
    content = (await upload.read()).decode("utf-8-sig")
    return list(csv.DictReader(StringIO(content)))


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


async def read_workbook(upload: UploadFile) -> dict[str, list[dict[str, str]]]:
    try:
        workbook = load_workbook(BytesIO(await upload.read()), data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="Upload must be a valid .xlsx Excel workbook") from exc
    sheets: dict[str, list[dict[str, str]]] = {}
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            sheets[sheet_name.lower()] = []
            continue
        headers = [normalize_cell(value) for value in rows[0]]
        sheet_rows: list[dict[str, str]] = []
        for values in rows[1:]:
            if not any(value is not None and str(value).strip() for value in values):
                continue
            sheet_rows.append({header: normalize_cell(value) for header, value in zip(headers, values) if header})
        sheets[sheet_name.lower()] = sheet_rows
    return sheets


@router.post("/import/households", response_model=ImportPreview)
async def import_households(file: UploadFile = File(...), db: Session = Depends(get_db)) -> ImportPreview:
    rows = await read_csv(file)
    errors: list[ImportErrorRow] = []
    objects: list[Household] = []
    seen: set[str] = set()
    existing = {code for (code,) in db.query(Household.household_code).all()}
    for index, row in enumerate(rows, start=2):
        row_errors: list[str] = []
        code = (row.get("household_code") or "").strip()
        if not code:
            row_errors.append("household_code is required")
        if code in seen or code in existing:
            row_errors.append("household_code must be unique")
        status_value = (row.get("status") or "active").strip() or "active"
        if status_value not in {"active", "inactive"}:
            row_errors.append("status must be active or inactive")
        if row_errors:
            errors.append(ImportErrorRow(row=index, errors=row_errors))
            continue
        seen.add(code)
        objects.append(Household(household_code=code, location=row.get("location"), private_name=row.get("private_name"), private_phone=row.get("private_phone"), private_notes=row.get("private_notes"), status=status_value))
    if errors:
        return ImportPreview(valid=False, rows_valid=len(objects), errors=errors)
    db.add_all(objects)
    commit_or_conflict(db)
    return ImportPreview(valid=True, rows_valid=len(objects), errors=[])


@router.post("/import/donations", response_model=ImportPreview)
async def import_donations(file: UploadFile = File(...), db: Session = Depends(get_db)) -> ImportPreview:
    rows = await read_csv(file)
    errors: list[ImportErrorRow] = []
    objects: list[DonationReceived] = []
    for index, row in enumerate(rows, start=2):
        row_errors: list[str] = []
        amount = parse_decimal(row.get("amount", ""), "amount", row_errors)
        received_date = parse_date(row.get("received_date", ""), "received_date", row_errors)
        if row_errors:
            errors.append(ImportErrorRow(row=index, errors=row_errors))
            continue
        objects.append(DonationReceived(donor_name=row.get("donor_name"), donor_display_name=row.get("donor_display_name"), donor_country=row.get("donor_country"), is_public=(row.get("is_public", "").lower() in {"true", "1", "yes"}), amount=amount, currency="USD", received_date=received_date, notes=row.get("notes")))
    if errors:
        return ImportPreview(valid=False, rows_valid=len(objects), errors=errors)
    db.add_all(objects)
    commit_or_conflict(db)
    return ImportPreview(valid=True, rows_valid=len(objects), errors=[])


@router.post("/import/distributions", response_model=ImportPreview)
async def import_distributions(file: UploadFile = File(...), db: Session = Depends(get_db)) -> ImportPreview:
    rows = await read_csv(file)
    errors: list[ImportErrorRow] = []
    objects: list[Distribution] = []
    households = {code: hid for hid, code in db.query(Household.id, Household.household_code).all()}
    for index, row in enumerate(rows, start=2):
        row_errors: list[str] = []
        code = (row.get("household_code") or "").strip()
        if code not in households:
            row_errors.append("household_code must match an existing household")
        amount = parse_decimal(row.get("amount", ""), "amount", row_errors)
        distribution_date = parse_date(row.get("distribution_date", ""), "distribution_date", row_errors)
        if row_errors:
            errors.append(ImportErrorRow(row=index, errors=row_errors))
            continue
        objects.append(Distribution(distribution_code=row.get("distribution_code") or None, household_id=households[code], amount=amount, currency="USD", distribution_date=distribution_date, assistance_type=row.get("assistance_type"), notes=row.get("notes")))
    if errors:
        return ImportPreview(valid=False, rows_valid=len(objects), errors=errors)
    db.add_all(objects)
    commit_or_conflict(db)
    return ImportPreview(valid=True, rows_valid=len(objects), errors=[])


@router.post("/import/workbook", response_model=ImportPreview)
async def import_workbook(file: UploadFile = File(...), db: Session = Depends(get_db)) -> ImportPreview:
    sheets = await read_workbook(file)
    household_rows = sheets.get("households")
    distribution_rows = sheets.get("distributions")
    if household_rows is None or distribution_rows is None:
        return ImportPreview(
            valid=False,
            rows_valid=0,
            errors=[ImportErrorRow(row=1, errors=["Workbook must include households and distributions sheets"])],
        )

    errors: list[ImportErrorRow] = []
    valid_households: list[dict[str, str]] = []
    valid_distributions: list[dict[str, str | Decimal | date | None]] = []
    seen_households: set[str] = set()
    seen_distribution_codes: set[str] = set()

    existing_households = {code: hid for hid, code in db.query(Household.id, Household.household_code).all()}
    for index, row in enumerate(household_rows, start=2):
        row_errors: list[str] = []
        code = (row.get("household_code") or "").strip()
        if not code:
            row_errors.append("household_code is required")
        if code in seen_households:
            row_errors.append("household_code appears more than once in households sheet")
        status_value = (row.get("status") or "active").strip() or "active"
        if status_value not in {"active", "inactive"}:
            row_errors.append("status must be active or inactive")
        if row_errors:
            errors.append(ImportErrorRow(row=index, errors=row_errors))
            continue
        seen_households.add(code)
        valid_households.append({**row, "household_code": code, "status": status_value})

    available_households = set(existing_households) | seen_households
    for index, row in enumerate(distribution_rows, start=2):
        row_errors: list[str] = []
        code = (row.get("household_code") or "").strip()
        distribution_code = (row.get("distribution_code") or "").strip()
        if distribution_code and distribution_code in seen_distribution_codes:
            row_errors.append("distribution_code appears more than once in distributions sheet")
        if distribution_code:
            seen_distribution_codes.add(distribution_code)
        if not code or code not in available_households:
            row_errors.append("household_code must match an existing or workbook household")
        amount = parse_decimal(row.get("amount", ""), "amount", row_errors)
        distribution_date = parse_date(row.get("distribution_date", ""), "distribution_date", row_errors)
        if row_errors:
            errors.append(ImportErrorRow(row=index, errors=row_errors))
            continue
        valid_distributions.append(
            {
                **row,
                "household_code": code,
                "distribution_code": distribution_code or None,
                "amount": amount,
                "distribution_date": distribution_date,
            }
        )

    if errors:
        return ImportPreview(valid=False, rows_valid=len(valid_households) + len(valid_distributions), errors=errors)

    for row in valid_households:
        household = db.query(Household).filter(Household.household_code == row["household_code"]).first()
        if not household:
            household = Household(household_code=row["household_code"])
            db.add(household)
            db.flush()
        household.location = row.get("location") or household.location
        household.private_name = row.get("private_name") or household.private_name
        household.private_phone = row.get("private_phone") or household.private_phone
        household.private_notes = row.get("private_notes") or household.private_notes
        household.status = row.get("status") or household.status
        existing_households[household.household_code] = household.id

    created_or_updated_distributions = 0
    for row in valid_distributions:
        household_id = existing_households[str(row["household_code"])]
        distribution = None
        distribution_code = row.get("distribution_code")
        if distribution_code:
            distribution = db.query(Distribution).filter(Distribution.distribution_code == distribution_code).first()
        if distribution is None and not distribution_code:
            distribution = (
                db.query(Distribution)
                .filter(
                    Distribution.household_id == household_id,
                    Distribution.amount == row["amount"],
                    Distribution.distribution_date == row["distribution_date"],
                    Distribution.assistance_type == (row.get("assistance_type") or None),
                )
                .first()
            )
        if not distribution:
            distribution = Distribution(household_id=household_id)
            db.add(distribution)
        distribution.distribution_code = distribution_code or distribution.distribution_code
        distribution.household_id = household_id
        distribution.amount = row["amount"]
        distribution.currency = "USD"
        distribution.distribution_date = row["distribution_date"]
        distribution.assistance_type = row.get("assistance_type") or None
        distribution.notes = row.get("notes") or distribution.notes
        created_or_updated_distributions += 1

    commit_or_conflict(db)
    return ImportPreview(valid=True, rows_valid=len(valid_households) + created_or_updated_distributions, errors=[])
